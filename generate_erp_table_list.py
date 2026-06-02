#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从思方云2 ERP Excel 生成表名清单 Markdown（格式对齐 MES V1.2）。"""

from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl

EXCEL_PATH = Path(__file__).resolve().parent / "思方云2_ERP数据库字段V3.xlsx"
OUTPUT_PATH = Path(__file__).resolve().parent / "中络项目ERP 系统表名清单V1.0.md"
DB_NAME = "思方云2 ERP"

# 模块分组（顺序与 数据库命名规范 一致）
MODULE_RULES: List[Tuple[str, str, callable]] = [
    ("2.1", "基础模块", lambda t: t.startswith("T_")),
    ("2.2", "物料模块", lambda t: t.startswith("M_")),
    ("2.3", "成品模块", lambda t: t.startswith("FGI_")),
    ("2.4", "工程模块", lambda t: t.startswith("E_")),
    ("2.5", "销售模块", lambda t: t.startswith("S_")),
    ("2.6", "生产模块", lambda t: t.startswith("P_")),
    ("2.7", "品质模块", lambda t: t.startswith("Q_")),
    ("2.8", "财务模块", lambda t: t.startswith("F_")),
    ("2.9", "成本模块", lambda t: t.startswith("C_")),
    ("2.10", "OA办公模块", lambda t: t.startswith("OA_")),
    ("2.11", "设备管理", lambda t: t.startswith("EQ_") or t.startswith("PM_")),
    ("2.12", "APS排程", lambda t: t in {"STEP", "EQUIPMENT", "EQUIPMENT_GROUP"}),
    ("2.13", "工作流引擎", lambda t: t.startswith("JBPM4")),
    ("2.15", "其它", lambda t: True),
]

CUSTOM_TABLES: Set[str] = set()
TABLE_PK: Dict[str, str] = defaultdict(lambda: "recId")


def _clean(val: Any) -> Optional[str]:
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def _is_valid_table_name(name: str) -> bool:
    if not name or name in ("表名", "表", "S_ShipmentRevokeItem"):
        return False
    invalid_markers = ("--", " as ", "\n", "select ", " from ")
    lower = name.lower()
    return not any(m in lower for m in invalid_markers)


def _is_valid_related_table(related: str) -> bool:
    upper = related.upper()
    sql_tokens = ("SELECT ", " FROM ", " JOIN ", " WHERE ", " GROUP BY ", " ORDER BY ")
    return not any(tok in upper for tok in sql_tokens)


def _format_relation(source_table: str, source_field: str, related: str) -> Optional[str]:
    related = related.strip()
    if not _is_valid_related_table(related):
        return None
    if "=>" in related:
        left, right = related.split("=>", 1)
        left = left.strip()
        right = right.strip()
        if "." in left:
            rt, rf = left.rsplit(".", 1)
            return f"{source_table}.{source_field} = {rt}.{right or rf}"
        return f"{source_table}.{source_field} = {left}.{right}"
    if "." in related:
        rt, rf = related.rsplit(".", 1)
        return f"{source_table}.{source_field} = {rt}.{rf}"
    pk = TABLE_PK.get(related, "recId")
    return f"{source_table}.{source_field} = {related}.{pk}"


def _table_chinese_name(table: str, meta: Dict[str, Any]) -> str:
    for key in ("desc_from_name_sheet", "comment_from_fields", "comment_from_cost", "comment_from_aps"):
        val = _clean(meta.get(key))
        if val and val not in ("描述", "字段说明", "表说明", "字段注释"):
            return val
    return table


def _business_meaning(table: str, meta: Dict[str, Any]) -> str:
    for key in ("desc_from_name_sheet", "comment_from_fields", "comment_from_cost", "comment_from_aps"):
        val = _clean(meta.get(key))
        if val and val not in ("描述", "字段说明", "表说明", "字段注释"):
            return val
    cn = _table_chinese_name(table, meta)
    if cn != table:
        return cn
    return f"ERP 系统 {table} 业务数据表"


def _classify_table(table: str) -> Tuple[str, str]:
    if table in CUSTOM_TABLES:
        return "2.14", "自定义表"
    for sec_id, sec_name, pred in MODULE_RULES:
        if sec_id == "2.15":
            return sec_id, sec_name
        if pred(table):
            return sec_id, sec_name
    return "2.15", "其它"


def parse_fields_sheet(wb) -> Dict[str, Dict[str, Any]]:
    ws = wb["字段注释"]
    tables: Dict[str, Dict[str, Any]] = {}
    current: Optional[str] = None

    for row in ws.iter_rows(values_only=True):
        if row[0] == "表名":
            continue
        if row[0]:
            current = str(row[0]).strip()
            if not _is_valid_table_name(current):
                current = None
                continue
            tables[current] = {
                "source_sheets": {"字段注释"},
                "comment_from_fields": _clean(row[3]),
                "fields": [],
                "relations": [],
            }
            continue
        if not current or not row[2]:
            continue
        field = str(row[2]).strip()
        pk = _clean(row[7])
        if pk:
            TABLE_PK[current] = field
        rel = _clean(row[5])
        entry = {
            "field": field,
            "comment": _clean(row[3]),
            "related_table": rel,
        }
        tables[current]["fields"].append(entry)
        if rel:
            formatted = _format_relation(current, field, rel)
            if formatted:
                tables[current]["relations"].append(formatted)
    return tables


def parse_name_sheet(wb, tables: Dict[str, Dict[str, Any]]) -> None:
    ws = wb["表名"]
    for row in ws.iter_rows(values_only=True):
        name = _clean(row[0])
        if not name or not _is_valid_table_name(name):
            continue
        desc = _clean(row[1])
        if desc == "描述":
            continue
        if name not in tables:
            tables[name] = {
                "source_sheets": set(),
                "comment_from_fields": None,
                "fields": [],
                "relations": [],
            }
        tables[name]["source_sheets"].add("表名")
        if desc:
            tables[name]["desc_from_name_sheet"] = desc


def parse_cost_sheet(wb, tables: Dict[str, Dict[str, Any]]) -> None:
    ws = wb["成本"]
    current: Optional[str] = None
    for row in ws.iter_rows(values_only=True):
        if row[0] == "表名":
            continue
        if row[0]:
            current = str(row[0]).strip()
            if not _is_valid_table_name(current):
                current = None
                continue
            if current not in tables:
                tables[current] = {
                    "source_sheets": set(),
                    "comment_from_fields": None,
                    "fields": [],
                    "relations": [],
                }
            tables[current]["source_sheets"].add("成本")
            tables[current]["comment_from_cost"] = _clean(row[3])
            continue
        if not current or not row[2]:
            continue
        field = str(row[2]).strip()
        fk = _clean(row[4])
        pk = _clean(row[6])
        if pk:
            TABLE_PK[current] = field
        if fk:
            rel = _format_relation(current, field, fk)
            if rel and rel not in tables[current]["relations"]:
                tables[current]["relations"].append(rel)


def parse_custom_sheet(wb, tables: Dict[str, Dict[str, Any]]) -> None:
    global CUSTOM_TABLES
    ws = wb["自定义表"]
    for row in ws.iter_rows(values_only=True):
        name = _clean(row[0])
        if not name or not _is_valid_table_name(name):
            continue
        CUSTOM_TABLES.add(name)
        desc = _clean(row[1])
        if name not in tables:
            tables[name] = {
                "source_sheets": set(),
                "comment_from_fields": None,
                "fields": [],
                "relations": [],
            }
        tables[name]["source_sheets"].add("自定义表")
        if desc:
            tables[name]["desc_from_name_sheet"] = desc


def parse_aps_sheet(wb, tables: Dict[str, Dict[str, Any]]) -> None:
    ws = wb["APS数据库"]
    current: Optional[str] = None
    for row in ws.iter_rows(values_only=True):
        if row[0] == "表名":
            continue
        if row[0]:
            current = str(row[0]).strip()
            if not _is_valid_table_name(current):
                current = None
                continue
            if current not in tables:
                tables[current] = {
                    "source_sheets": set(),
                    "comment_from_fields": None,
                    "fields": [],
                    "relations": [],
                }
            tables[current]["source_sheets"].add("APS数据库")
            tables[current]["comment_from_aps"] = _clean(row[3])
            continue
        if not current or not row[2]:
            continue
        field = str(row[2]).strip()
        mapped = _clean(row[4])
        if mapped:
            rel = _format_relation(current, field, mapped)
            if rel and rel not in tables[current]["relations"]:
                tables[current]["relations"].append(rel)


def build_summary_sections(tables: Dict[str, Dict[str, Any]]) -> str:
    summary_map = [
        ("基础模块核心关联", ["T_"]),
        ("物料模块核心关联", ["M_"]),
        ("成品模块核心关联", ["FGI_"]),
        ("工程模块核心关联", ["E_"]),
        ("销售模块核心关联", ["S_"]),
        ("生产模块核心关联", ["P_"]),
        ("品质模块核心关联", ["Q_"]),
        ("财务模块核心关联", ["F_"]),
        ("成本模块核心关联", ["C_"]),
    ]
    lines = ["## 三、核心关联关系汇总", ""]
    sec_no = 1
    for title, prefixes in summary_map:
        rels: List[str] = []
        seen: Set[str] = set()
        for table, meta in sorted(tables.items()):
            if not any(table.startswith(p) for p in prefixes):
                continue
            for rel in meta.get("relations", []):
                if rel not in seen:
                    seen.add(rel)
                    rels.append(rel)
                if len(rels) >= 8:
                    break
            if len(rels) >= 8:
                break
        if not rels:
            continue
        lines.append(f"### 3.{sec_no} {title}")
        lines.append("")
        lines.append("| 序号 | 关联关系 |")
        lines.append("|------|----------|")
        for i, rel in enumerate(rels[:8], 1):
            lines.append(f"| {i} | {rel} |")
        lines.append("")
        sec_no += 1
    return "\n".join(lines)


def render_table_block(idx: int, table: str, meta: Dict[str, Any]) -> str:
    cn = _table_chinese_name(table, meta)
    meaning = _business_meaning(table, meta)
    rels = meta.get("relations") or []
    lines = [
        f"#### {idx} {cn} ({table})",
        "",
        f"- **业务含义**：{meaning}",
        f"- **所属数据库**：{DB_NAME}",
    ]
    if rels:
        lines.append("-关联关系：")
        lines.extend(f"- {r}" for r in rels)
    else:
        lines.append("-关联关系：无")
    lines.extend(["", "---", ""])
    return "\n".join(lines)


def generate() -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    tables = parse_fields_sheet(wb)
    parse_name_sheet(wb, tables)
    parse_cost_sheet(wb, tables)
    parse_custom_sheet(wb, tables)
    parse_aps_sheet(wb, tables)
    wb.close()

    sections: Dict[str, Dict[str, Any]] = defaultdict(lambda: {"name": "", "tables": []})
    for table in sorted(tables.keys(), key=lambda x: (x.lower(), x)):
        sec_id, sec_name = _classify_table(table)
        sections[sec_id]["name"] = sec_name
        sections[sec_id]["tables"].append(table)

    out: List[str] = [
        "# 数据库结构参考说明文档（关联关系版）",
        "",
        "",
        "## 二、核心数据表结构及关联关系",
        "",
    ]

    global_idx = 1
    ordered_sec_ids = sorted(sections.keys(), key=lambda x: float(x.split()[0].replace("2.", "")))
    for sec_id in ordered_sec_ids:
        sec = sections[sec_id]
        sheet_sources = set()
        for t in sec["tables"]:
            sheet_sources.update(tables[t].get("source_sheets", set()))
        source_note = "、".join(sorted(sheet_sources)) if sheet_sources else "字段注释"
        out.append(f"### {sec_id} {sec['name']}")
        out.append("")
        out.append(f"> 本章节数据来源于 Excel 工作表：`{source_note}`")
        out.append("")
        for table in sec["tables"]:
            out.append(render_table_block(global_idx, table, tables[table]))
            global_idx += 1

    out.append(build_summary_sections(tables))
    out.append("")
    out.append("---")
    out.append("")

    content = "\n".join(out)
    OUTPUT_PATH.write_text(content, encoding="utf-8")
    print(f"Generated {OUTPUT_PATH}")
    print(f"Total tables: {global_idx - 1}")
    print(f"Sections: {len(sections)}")


if __name__ == "__main__":
    generate()
