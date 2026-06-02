#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从思方云2 ERP Excel 生成数据库表结构 Markdown（格式对齐 MES V1.2）。"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl

import generate_erp_table_list as erp

EXCEL_PATH = erp.EXCEL_PATH
OUTPUT_PATH = Path(__file__).resolve().parent / "中络项目ERP 系统数据库表结构V1.0.md"
DB_NAME = erp.DB_NAME
DOC_VERSION = "V1.0"

NULLABLE_MARKS = {"√", "Y", "y", "是", "YES", "yes", "1", True}

SQL_TYPE_MAP = {
    "nvarchar": "string",
    "varchar": "string",
    "char": "string",
    "nchar": "string",
    "text": "string",
    "ntext": "string",
    "xml": "string",
    "uniqueidentifier": "string",
    "int": "int",
    "bigint": "long",
    "smallint": "int",
    "tinyint": "int",
    "numeric": "decimal",
    "decimal": "decimal",
    "money": "decimal",
    "smallmoney": "decimal",
    "float": "double",
    "real": "double",
    "bit": "bool",
    "datetime": "DateTime",
    "datetime2": "DateTime",
    "smalldatetime": "DateTime",
    "date": "DateTime",
    "time": "TimeSpan",
    "image": "byte[]",
    "varbinary": "byte[]",
    "binary": "byte[]",
}


def _is_nullable(val: Any) -> bool:
    if val is None:
        return True
    if isinstance(val, bool):
        return val
    s = str(val).strip()
    if not s:
        return True
    return s in NULLABLE_MARKS


def _map_field_type(sql_type: Optional[str], nullable: bool) -> str:
    if not sql_type:
        return "string"
    base = SQL_TYPE_MAP.get(sql_type.lower().strip(), "string")
    if nullable and base not in ("string", "byte[]") and not base.endswith("?"):
        return f"{base}?"
    return base


def _default_display(val: Any) -> str:
    cleaned = erp._clean(val)
    if not cleaned:
        return "-"
    return cleaned


def _field_desc(
    comment: Optional[str],
    value_desc: Optional[str],
    related: Optional[str],
) -> str:
    parts: List[str] = []
    if comment:
        parts.append(comment)
    if value_desc and value_desc not in parts:
        parts.append(value_desc)
    desc = "；".join(parts) if parts else ""
    if related and erp._is_valid_related_table(related):
        if "." in related:
            desc = f"{desc}，对应{related}" if desc else f"对应{related}"
        else:
            pk = erp.TABLE_PK.get(related, "recId")
            suffix = f"，对应{related}.{pk}"
            desc = f"{desc}{suffix}" if desc else f"对应{related}.{pk}"
    return desc or "-"


def _upsert_field(
    tables: Dict[str, Dict[str, Any]],
    table: str,
    field: str,
    *,
    comment: Optional[str] = None,
    value_desc: Optional[str] = None,
    related: Optional[str] = None,
    sql_type: Optional[str] = None,
    nullable: Optional[bool] = None,
    default_val: Any = None,
    is_pk: bool = False,
    source_sheet: str,
) -> None:
    meta = tables[table]
    meta["source_sheets"].add(source_sheet)
    by_name: Dict[str, Dict[str, Any]] = meta.setdefault("_field_map", {})
    entry = by_name.get(field)
    if not entry:
        entry = {
            "field": field,
            "comment": None,
            "value_desc": None,
            "related_table": None,
            "sql_type": None,
            "nullable": True,
            "default": None,
            "is_pk": False,
        }
        by_name[field] = entry
        meta["fields"].append(entry)
    if comment:
        entry["comment"] = comment
    if value_desc:
        entry["value_desc"] = value_desc
    if related:
        entry["related_table"] = related
        formatted = erp._format_relation(table, field, related)
        if formatted and formatted not in meta["relations"]:
            meta["relations"].append(formatted)
    if sql_type:
        entry["sql_type"] = sql_type
    if nullable is not None:
        entry["nullable"] = nullable
    if default_val is not None and erp._clean(default_val):
        entry["default"] = erp._clean(default_val)
    if is_pk:
        entry["is_pk"] = True
        entry["nullable"] = False
        erp.TABLE_PK[table] = field


def _new_table_meta() -> Dict[str, Any]:
    return {
        "source_sheets": set(),
        "comment_from_fields": None,
        "fields": [],
        "relations": [],
        "_field_map": {},
    }


def parse_fields_sheet_full(wb) -> Dict[str, Dict[str, Any]]:
    ws = wb["字段注释"]
    tables: Dict[str, Dict[str, Any]] = {}
    current: Optional[str] = None

    for row in ws.iter_rows(values_only=True):
        if row[0] == "表名":
            continue
        if row[0]:
            current = str(row[0]).strip()
            if not erp._is_valid_table_name(current):
                current = None
                continue
            tables[current] = _new_table_meta()
            tables[current]["source_sheets"].add("字段注释")
            tables[current]["comment_from_fields"] = erp._clean(row[3])
            continue
        if not current or not row[2]:
            continue
        field = str(row[2]).strip()
        comment = erp._clean(row[3])
        is_pk = (
            erp._clean(row[7]) in NULLABLE_MARKS
            or (field.lower() == "recid" and comment and "主键" in comment)
        )
        _upsert_field(
            tables,
            current,
            field,
            comment=comment,
            value_desc=erp._clean(row[4]),
            related=erp._clean(row[5]),
            sql_type=erp._clean(row[8]),
            nullable=_is_nullable(row[12] if len(row) > 12 else None),
            default_val=row[13] if len(row) > 13 else None,
            is_pk=is_pk,
            source_sheet="字段注释",
        )
    return tables


def parse_cost_sheet_full(wb, tables: Dict[str, Dict[str, Any]]) -> None:
    ws = wb["成本"]
    current: Optional[str] = None
    for row in ws.iter_rows(values_only=True):
        if row[0] == "表名":
            continue
        if row[0]:
            current = str(row[0]).strip()
            if not erp._is_valid_table_name(current):
                current = None
                continue
            if current not in tables:
                tables[current] = _new_table_meta()
            tables[current]["source_sheets"].add("成本")
            if erp._clean(row[3]):
                tables[current]["comment_from_cost"] = erp._clean(row[3])
            continue
        if not current or not row[2]:
            continue
        field = str(row[2]).strip()
        is_pk = erp._clean(row[6]) in NULLABLE_MARKS
        _upsert_field(
            tables,
            current,
            field,
            comment=erp._clean(row[3]),
            related=erp._clean(row[4]),
            sql_type=erp._clean(row[7]),
            nullable=_is_nullable(row[11] if len(row) > 11 else None),
            is_pk=is_pk,
            source_sheet="成本",
        )


def parse_aps_sheet_full(wb, tables: Dict[str, Dict[str, Any]]) -> None:
    ws = wb["APS数据库"]
    current: Optional[str] = None
    for row in ws.iter_rows(values_only=True):
        if row[0] == "表名":
            continue
        if row[0]:
            current = str(row[0]).strip()
            if not erp._is_valid_table_name(current):
                current = None
                continue
            if current not in tables:
                tables[current] = _new_table_meta()
            tables[current]["source_sheets"].add("APS数据库")
            if erp._clean(row[3]):
                tables[current]["comment_from_aps"] = erp._clean(row[3])
            continue
        if not current or not row[2]:
            continue
        field = str(row[2]).strip()
        _upsert_field(
            tables,
            current,
            field,
            comment=erp._clean(row[3]),
            related=erp._clean(row[4]),
            sql_type=erp._clean(row[5]),
            nullable=_is_nullable(row[9] if len(row) > 9 else None),
            source_sheet="APS数据库",
        )


def render_field_row(entry: Dict[str, Any]) -> str:
    nullable = entry.get("nullable", True)
    if entry.get("is_pk"):
        nullable = False
    null_text = "是" if nullable else "否"
    ftype = _map_field_type(entry.get("sql_type"), nullable)
    desc = _field_desc(
        entry.get("comment"),
        entry.get("value_desc"),
        entry.get("related_table"),
    )
    if entry.get("is_pk") and desc == "-":
        desc = "主键"
    elif entry.get("is_pk") and "主键" not in desc:
        desc = f"{desc}（主键）" if desc != "-" else "主键"
    return (
        f"| {entry['field']} | {ftype} | {null_text} | "
        f"{_default_display(entry.get('default'))} | {desc} |"
    )


def render_table_block(idx: int, table: str, meta: Dict[str, Any]) -> List[str]:
    cn = erp._table_chinese_name(table, meta)
    meaning = erp._business_meaning(table, meta)
    rels = meta.get("relations") or []
    fields: List[Dict[str, Any]] = meta.get("fields") or []

    lines = [
        f"#### {idx} {cn} ( {table} )",
        f"- **业务含义**：{meaning}",
        f"- **所属数据库**：{DB_NAME}",
        "| 字段名 | 字段类型 | 是否为空 | 默认值 | 说明 |",
        "|--------|----------|----------|--------|------|",
    ]
    if fields:
        for entry in fields:
            lines.append(render_field_row(entry))
    else:
        lines.append("| — | — | — | — | 数据字典中暂无字段明细 |")
    if rels:
        lines.append("- **关联关系**：")
        lines.extend(f"  - {r}" for r in rels)
    else:
        lines.append("- **关联关系**：无")
    lines.extend(["", "---", ""])
    return lines


def build_header() -> List[str]:
    return [
        '<h1 align="center">数据库结构参考说明文档</h1>',
        "",
        "> 本文档根据数据字典生成，为使用人员提供数据库表结构参考，帮助快速熟悉表定义及字段含义",
        "",
        "## 变更记录",
        "",
        "| 序号 | 变更内容 | 变更时间 | 变更人 | 备注 |",
        "|------|----------|----------|--------|------|",
        f"| 1 | {DOC_VERSION} | 待填写 | 自动生成脚本 | 初始版本 |",
        "|  |  |  |  |  |",
        "",
        "## 一、文档概述",
        "",
        "### 1.1 文档目的",
        "",
        "本文档旨在为使用人员提供数据库结构的全面参考，使使用人员能够：",
        "- 快速理解各数据表的用途及字段含义",
        "- 掌握各表所属业务域与数据库来源",
        "- 准确编写查询语句进行数据提取",
        "- 快速完成系统数据字典查阅",
        "",
        "---",
        "",
        "## 二、核心数据表结构",
        "",
    ]


def build_relation_section(tables: Dict[str, Dict[str, Any]]) -> List[str]:
    lines = [
        "## 三、表关系说明",
        "",
        "### 3.1 核心关联关系摘录",
        "",
        "以下摘录各业务模块中部分主外键关联，完整关联见第二章各表「关联关系」小节。",
        "",
    ]
    summary_map = [
        ("基础模块", ["T_"]),
        ("物料模块", ["M_"]),
        ("成品模块", ["FGI_"]),
        ("工程模块", ["E_"]),
        ("销售模块", ["S_"]),
        ("生产模块", ["P_"]),
        ("品质模块", ["Q_"]),
        ("财务模块", ["F_"]),
        ("成本模块", ["C_"]),
    ]
    sec_no = 1
    for title, prefixes in summary_map:
        rels: List[str] = []
        seen: Set[str] = set()
        for table, meta in sorted(tables.items()):
            if not any(table.startswith(p) for p in prefixes):
                continue
            for rel in meta.get("relations", []):
                if rel not in seen:
                    seen.add(rel)
                    rels.append(rel)
                if len(rels) >= 12:
                    break
            if len(rels) >= 12:
                break
        if not rels:
            continue
        lines.append(f"### 3.{sec_no} {title}")
        lines.append("")
        lines.append("| 序号 | 关联关系 |")
        lines.append("|------|----------|")
        for i, rel in enumerate(rels, 1):
            lines.append(f"| {i} | {rel} |")
        lines.append("")
        sec_no += 1
    if sec_no == 1:
        lines.extend(
            [
                "### 3.1 表关系占位说明",
                "",
                "当前模块未解析到可展示的关联关系，请参见第二章各表「关联关系」小节。",
                "",
            ]
        )
    lines.extend(
        [
            "### 3.2 推荐补充模板",
            "",
            "| 主表 | 外键表 | 关系类型 | 说明 |",
            "|------|--------|----------|------|",
            "| 待补充 | 待补充 | 待补充 | 待补充 |",
            "",
        ]
    )
    return lines


def build_index_section(tables: Dict[str, Dict[str, Any]], ordered_tables: List[str]) -> List[str]:
    lines = [
        "## 四、数据字典速查",
        "",
        "### 4.1 表清单",
        "",
        "| 序号 | 表名 | 表注释 | 所属数据库 | 字段数量 |",
        "|------|------|----------|------------|----------|",
    ]
    for i, table in enumerate(ordered_tables, 1):
        meta = tables[table]
        cn = erp._table_chinese_name(table, meta)
        field_count = len(meta.get("fields") or [])
        lines.append(f"| {i} | {table} | {cn} | {DB_NAME} | {field_count} |")
    lines.extend(["", "---", ""])
    return lines


def generate() -> None:
    erp.CUSTOM_TABLES = set()
    erp.TABLE_PK = defaultdict(lambda: "recId")

    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    tables = parse_fields_sheet_full(wb)
    erp.parse_name_sheet(wb, tables)
    parse_cost_sheet_full(wb, tables)
    erp.parse_custom_sheet(wb, tables)
    parse_aps_sheet_full(wb, tables)
    wb.close()

    for meta in tables.values():
        meta.pop("_field_map", None)

    sections: Dict[str, Dict[str, Any]] = defaultdict(lambda: {"name": "", "tables": []})
    for table in sorted(tables.keys(), key=lambda x: (x.lower(), x)):
        sec_id, sec_name = erp._classify_table(table)
        sections[sec_id]["name"] = sec_name
        sections[sec_id]["tables"].append(table)

    out: List[str] = build_header()
    global_idx = 1
    ordered_tables: List[str] = []
    ordered_sec_ids = sorted(sections.keys(), key=lambda x: float(x.replace("2.", "")))

    for sec_id in ordered_sec_ids:
        sec = sections[sec_id]
        sheet_sources: Set[str] = set()
        for t in sec["tables"]:
            sheet_sources.update(tables[t].get("source_sheets", set()))
        source_note = "、".join(sorted(sheet_sources)) if sheet_sources else "字段注释"
        out.append(f"### {sec_id} {sec['name']}")
        out.append("")
        out.append(f"> 本章节数据来源于 Excel 工作表：`{source_note}`")
        out.append("")
        for table in sec["tables"]:
            out.extend(render_table_block(global_idx, table, tables[table]))
            ordered_tables.append(table)
            global_idx += 1

    out.extend(build_relation_section(tables))
    out.extend(build_index_section(tables, ordered_tables))

    content = "\n".join(out)
    OUTPUT_PATH.write_text(content, encoding="utf-8")
    total_fields = sum(len(tables[t].get("fields") or []) for t in tables)
    print(f"Generated {OUTPUT_PATH}")
    print(f"Total tables: {len(tables)}")
    print(f"Total fields: {total_fields}")
    print(f"Sections: {len(sections)}")


if __name__ == "__main__":
    generate()
